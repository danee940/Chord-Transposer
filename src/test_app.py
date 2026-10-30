"""Tests for the Flask web layer in app.py."""
# pylint: disable=redefined-outer-name,missing-function-docstring,protected-access

import json
import re
from collections.abc import Iterator
from datetime import date
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any, NoReturn
from urllib.parse import quote_plus, unquote

import pytest
from docx import Document
from flask.testing import FlaskClient
from pypdf import PdfReader, PdfWriter

import app as app_module
import seo
from app import app

DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _one_page_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture
def client() -> Iterator[FlaskClient]:
    app.config["TESTING"] = True
    with app.test_client() as test_client:
        yield test_client


def _chord_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("C G Am F")
    document.add_paragraph("These are the lyrics")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(data_bytes: bytes, filename: str = "song.docx") -> dict[str, object]:
    return {"file": (BytesIO(data_bytes), filename)}


def test_index_lists_keys(client: FlaskClient) -> None:
    response = client.get("/")
    assert response.status_code == 200
    assert b"Chord Transposer" in response.data


def test_index_shows_docx_formatting_moat_badge(client: FlaskClient) -> None:
    body = client.get("/").get_data(as_text=True)
    assert "Keeps your Word (.docx) formatting intact" in body


def test_index_links_versioned_stylesheet(client: FlaskClient) -> None:
    response = client.get("/")
    body = response.get_data(as_text=True)
    assert "cdn.tailwindcss.com" not in body
    assert f"/static/tailwind.css?v={app_module.CSS_VERSION}" in body


def test_index_html_is_not_cached(client: FlaskClient) -> None:
    response = client.get("/")
    assert response.headers["Cache-Control"] == "no-cache"


def test_static_css_is_cached_immutably(client: FlaskClient) -> None:
    response = client.get(f"/static/tailwind.css?v={app_module.CSS_VERSION}")
    assert response.status_code == 200
    cache_control = response.headers["Cache-Control"]
    assert "immutable" in cache_control
    assert f"max-age={app_module.STATIC_CSS_MAX_AGE}" in cache_control


def test_css_version_is_stable_hash() -> None:
    assert app_module.CSS_VERSION == app_module._compute_css_version()
    assert len(app_module.CSS_VERSION) == 12


def test_css_version_falls_back_when_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    def _raise(_self: Path) -> NoReturn:
        raise OSError("missing")

    monkeypatch.setattr(app_module.Path, "read_bytes", _raise)
    assert app_module._compute_css_version() == "dev"


def test_index_references_hashed_js_bundle(client: FlaskClient) -> None:
    response = client.get("/")
    body = response.get_data(as_text=True)
    assert f'src="{app_module.JS_BUNDLE}"' in body
    assert app_module.JS_BUNDLE.startswith("/static/js/")


def test_js_bundle_reads_hashed_entry_from_manifest(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        app_module.Path,
        "read_text",
        lambda _self: '{"main.ts": {"file": "main.abc123.js"}}',
    )
    assert app_module._resolve_js_bundle() == "/static/js/main.abc123.js"


def test_js_bundle_falls_back_when_manifest_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    def _raise(_self: Path) -> NoReturn:
        raise OSError("missing")

    monkeypatch.setattr(app_module.Path, "read_text", _raise)
    assert app_module._resolve_js_bundle() == app_module.JS_BUNDLE_FALLBACK


def test_js_bundle_falls_back_when_entry_malformed(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module.Path, "read_text", lambda _self: "{}")
    assert app_module._resolve_js_bundle() == app_module.JS_BUNDLE_FALLBACK


def test_health(client: FlaskClient) -> None:
    response = client.get("/health")
    assert response.status_code == 200
    assert response.get_json() == {"status": "ok"}


def test_robots_txt(client: FlaskClient) -> None:
    response = client.get("/robots.txt")
    assert response.status_code == 200
    assert response.mimetype == "text/plain"
    assert "Sitemap:" in response.get_data(as_text=True)


def test_og_image_png(client: FlaskClient) -> None:
    response = client.get("/og-image.png")
    assert response.status_code == 200
    assert response.mimetype == "image/png"
    assert response.get_data().startswith(b"\x89PNG")


def test_favicon_svg(client: FlaskClient) -> None:
    response = client.get("/favicon.svg")
    assert response.status_code == 200
    assert response.mimetype == "image/svg+xml"
    assert response.get_data(as_text=True).lstrip().startswith("<svg")


def test_favicon_ico(client: FlaskClient) -> None:
    response = client.get("/favicon.ico")
    assert response.status_code == 200
    assert response.mimetype == "image/vnd.microsoft.icon"
    assert response.data.startswith(b"\x00\x00\x01\x00")


def test_sitemap_xml(client: FlaskClient) -> None:
    response = client.get("/sitemap.xml")
    assert response.status_code == 200
    assert response.mimetype == "application/xml"
    assert "<urlset" in response.get_data(as_text=True)


def test_transpose_missing_file(client: FlaskClient) -> None:
    response = client.post(
        "/transpose",
        data={"current_key": "C", "target_key": "D"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "choose a .docx" in response.get_json()["error"]


def test_transpose_wrong_extension(client: FlaskClient) -> None:
    data = _upload(b"not a docx", filename="song.txt")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Only .docx" in response.get_json()["error"]


def test_transpose_missing_keys(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "", "target_key": ""})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "keys are required" in response.get_json()["error"]


def test_transpose_unsupported_format(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "C", "target_key": "D", "format": "txt"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Unsupported output format" in response.get_json()["error"]


def test_transpose_empty_file(client: FlaskClient) -> None:
    data = _upload(b"", filename="empty.docx")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "empty" in response.get_json()["error"]


def test_transpose_invalid_key(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "X", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "not a valid key" in response.get_json()["error"]


def test_transpose_invalid_docx(client: FlaskClient) -> None:
    data = _upload(b"this is not a real docx", filename="bad.docx")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "valid .docx" in response.get_json()["error"]


def test_transpose_docx_success(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")

    assert response.status_code == 200
    assert response.mimetype == DOCX_MIMETYPE
    assert "song_D.docx" in response.headers["Content-Disposition"]
    assert unquote(response.headers["X-Transpose-From"]) == "C"
    assert unquote(response.headers["X-Transpose-To"]) == "D"
    assert "C->D" in unquote(response.headers["X-Transpose-Changes"])


def test_transpose_preserves_non_ascii_filename(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes(), filename="Covered-Gizus_rövid_G.docx")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")

    assert response.status_code == 200
    disposition = response.headers["Content-Disposition"]
    assert "filename*=UTF-8''Covered-Gizus_r%C3%B6vid_G_D.docx" in disposition


def test_transpose_pdf_success(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "convert_docx_to_pdf", lambda docx_bytes: b"%PDF-1.4 fake")

    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "C", "target_key": "D", "format": "pdf"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")

    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert "song_D.pdf" in response.headers["Content-Disposition"]
    assert response.data == b"%PDF-1.4 fake"


def test_transpose_pdf_conversion_failure(
    client: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _raise(_docx_bytes: bytes) -> NoReturn:
        raise app_module.PdfConversionError("PDF conversion timed out.")

    monkeypatch.setattr(app_module, "convert_docx_to_pdf", _raise)

    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "C", "target_key": "D", "format": "pdf"})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")

    assert response.status_code == 503
    assert "timed out" in response.get_json()["error"]


def _jsonld_blocks(body: str) -> list[dict[str, Any]]:
    pattern = r'type="application/ld\+json">(.*?)</script>'
    return [json.loads(block) for block in re.findall(pattern, body, re.S)]


def _jsonld_types(body: str) -> set[str | None]:
    return {block.get("@type") for block in _jsonld_blocks(body)}


def _title(body: str) -> str:
    m = re.search(r"<title>(.*?)</title>", body, re.S)
    assert m is not None
    return m.group(1).strip()


def _canonical(body: str) -> str:
    m = re.search(r'<link rel="canonical" href="(.*?)"', body)
    assert m is not None
    return m.group(1)


def test_landing_pages_return_200_with_unique_title_and_canonical(client: FlaskClient) -> None:
    titles = {}
    canonicals = {}
    for page in seo.landing_pages():
        response = client.get(page["path"])
        assert response.status_code == 200, page["path"]
        body = response.get_data(as_text=True)
        titles.setdefault(_title(body), []).append(page["path"])
        canonicals.setdefault(_canonical(body), []).append(page["path"])
        assert _canonical(body) == f"{seo.SITE_URL}{page['path']}"

    landing_count = len(seo.landing_pages())
    assert len(titles) == landing_count
    assert len(canonicals) == landing_count


def test_landing_pages_include_only_breadcrumb_structured_data(client: FlaskClient) -> None:
    for page in seo.landing_pages():
        body = client.get(page["path"]).get_data(as_text=True)
        assert _jsonld_types(body) == {"BreadcrumbList"}, page["path"]


def test_landing_page_breadcrumbs_trail_back_to_home(client: FlaskClient) -> None:
    for page in seo.landing_pages():
        body = client.get(page["path"]).get_data(as_text=True)
        block = next(b for b in _jsonld_blocks(body) if b.get("@type") == "BreadcrumbList")
        items = block["itemListElement"]
        assert [item["position"] for item in items] == list(range(1, len(items) + 1))
        assert items[0]["item"] == f"{seo.SITE_URL}/"
        assert items[-1]["item"] == f"{seo.SITE_URL}{page['path']}"


def test_landing_page_title_differs_from_home(client: FlaskClient) -> None:
    home_title = _title(client.get("/").get_data(as_text=True))
    guitar_title = _title(client.get("/guitar-chord-transposer").get_data(as_text=True))
    assert guitar_title != home_title


def test_instrument_pages_exist(client: FlaskClient) -> None:
    for slug in ("guitar-chord-transposer", "ukulele-chord-transposer", "piano-chord-transposer"):
        assert client.get(f"/{slug}").status_code == 200


def test_chromatic_chart_page_renders_grid(client: FlaskClient) -> None:
    body = client.get("/chromatic-chart").get_data(as_text=True)
    assert "Chromatic transposition chart" in body
    assert "+7" in body


def test_chromatic_chart_page_shows_chart_before_transposer(client: FlaskClient) -> None:
    body = client.get("/chromatic-chart").get_data(as_text=True)
    assert body.index("Every key, shifted by semitone") < body.index('id="panel-paste"')


def test_chromatic_chart_page_has_reference_tables(client: FlaskClient) -> None:
    body = client.get("/chromatic-chart").get_data(as_text=True)
    assert "Semitones to interval names" in body
    assert "Perfect fifth" in body
    assert "Key signatures" in body
    assert "Relative minor" in body
    assert "F#m" in body


def test_home_has_no_data_instrument_attribute(client: FlaskClient) -> None:
    assert "data-instrument=" not in client.get("/").get_data(as_text=True)


def test_home_has_only_webapplication_jsonld(client: FlaskClient) -> None:
    body = client.get("/").get_data(as_text=True)
    assert _jsonld_types(body) == {"WebApplication"}


def test_no_page_emits_retired_rich_result_types(client: FlaskClient) -> None:
    for path in seo.sitemap_paths():
        body = client.get(path).get_data(as_text=True)
        assert not {"FAQPage", "HowTo"} & _jsonld_types(body), path


def test_faq_is_single_sourced(client: FlaskClient) -> None:
    body = client.get("/").get_data(as_text=True)
    assert body.count("<summary") == len(seo.FAQ_ITEMS)
    for item in seo.FAQ_ITEMS:
        assert escape(item["question"]) in body
        assert escape(item["answer"]) in body


def test_howto_steps_are_single_sourced(client: FlaskClient) -> None:
    body = client.get("/").get_data(as_text=True)
    for step in seo.HOWTO_STEPS:
        assert escape(step["text"]) in body


def test_home_has_no_capo_calculator(client: FlaskClient) -> None:
    body = client.get("/").get_data(as_text=True)
    assert 'id="capo-calculator"' not in body
    assert 'id="capo-key"' not in body


def test_guitar_page_has_capo_calculator(client: FlaskClient) -> None:
    body = client.get("/guitar-chord-transposer").get_data(as_text=True)
    assert 'id="capo-calculator"' in body
    assert 'id="capo-key"' in body


def test_instrument_pages_preselect_matching_instrument(client: FlaskClient) -> None:
    expected = {
        "guitar-chord-transposer": "guitar",
        "ukulele-chord-transposer": "ukulele",
        "piano-chord-transposer": "piano",
    }
    for slug, instrument in expected.items():
        body = client.get(f"/{slug}").get_data(as_text=True)
        assert f'data-instrument="{instrument}"' in body


def test_file_page_exists_with_upload_panel(client: FlaskClient) -> None:
    response = client.get("/file-chord-transposer")
    assert response.status_code == 200
    body = response.get_data(as_text=True)
    assert 'id="panel-upload"' in body
    assert 'id="file"' in body
    assert 'id="panel-paste"' not in body


def test_file_page_has_distinct_title_and_description(client: FlaskClient) -> None:
    body = client.get("/file-chord-transposer").get_data(as_text=True)
    home_body = client.get("/").get_data(as_text=True)
    assert _title(body) != _title(home_body)
    description_match = re.search(r'<meta name="description" content="(.*?)"', body)
    assert description_match is not None
    assert ".docx" in description_match.group(1)


def test_pages_have_page_specific_faq(client: FlaskClient) -> None:
    cases = {
        "/guitar-chord-transposer": seo.GUITAR_FAQ_ITEMS,
        "/ukulele-chord-transposer": seo.UKULELE_FAQ_ITEMS,
        "/piano-chord-transposer": seo.PIANO_FAQ_ITEMS,
        "/file-chord-transposer": seo.FILE_FAQ_ITEMS,
        "/chromatic-chart": seo.CHART_FAQ_ITEMS,
    }
    for path, faq_items in cases.items():
        body = client.get(path).get_data(as_text=True)
        assert body.count("<summary") == len(faq_items)
        for item in faq_items:
            assert escape(item["question"]) in body
            assert escape(item["answer"]) in body


def test_nav_lists_file_transposer_first(client: FlaskClient) -> None:
    sections = seo.nav_sections()
    assert sections[0]["links"][0]["path"] == "/file-chord-transposer"
    body = client.get("/").get_data(as_text=True)
    assert "/file-chord-transposer" in body


def test_nav_columns_balance_links_and_keep_all_sections() -> None:
    columns = seo.nav_columns()
    assert len(columns) == 2
    assert all(column for column in columns)
    flattened = [section for column in columns for section in column]
    assert sorted(s["heading"] for s in flattened) == sorted(
        s["heading"] for s in seo.nav_sections()
    )
    tall_column = max(columns, key=lambda column: sum(len(s["links"]) for s in column))
    assert any(s["heading"] == "By instrument" for s in tall_column)


def test_sitemap_lists_every_landing_url(client: FlaskClient) -> None:
    body = client.get("/sitemap.xml").get_data(as_text=True)
    for path in seo.sitemap_paths():
        assert f"<loc>{seo.SITE_URL}{path}</loc>" in body
    assert body.count("<url>") == len(seo.sitemap_paths())


def test_sitemap_gives_every_url_a_lastmod(client: FlaskClient) -> None:
    body = client.get("/sitemap.xml").get_data(as_text=True)
    for entry in seo.sitemap_entries():
        assert (
            f"<loc>{seo.SITE_URL}{entry['path']}</loc><lastmod>{entry['lastmod']}</lastmod>" in body
        )
    assert body.count("<lastmod>") == body.count("<url>")


def test_sitemap_lastmod_dates_are_iso_formatted(client: FlaskClient) -> None:
    body = client.get("/sitemap.xml").get_data(as_text=True)
    dates = re.findall(r"<lastmod>([^<]+)</lastmod>", body)
    assert dates
    for value in dates:
        assert date.fromisoformat(value).isoformat() == value


def test_sitemap_entries_cover_every_indexable_path() -> None:
    assert [entry["path"] for entry in seo.sitemap_entries()] == seo.sitemap_paths()
    assert all(entry["lastmod"] for entry in seo.sitemap_entries())


def test_every_landing_page_declares_a_lastmod() -> None:
    for page in seo.landing_pages():
        assert date.fromisoformat(page["lastmod"])


def test_robots_txt_points_to_sitemap(client: FlaskClient) -> None:
    body = client.get("/robots.txt").get_data(as_text=True)
    assert f"Sitemap: {seo.SITE_URL}/sitemap.xml" in body


RETIRED_KEY_PAIRS = [
    ("C", "D"),
    ("C", "G"),
    ("C", "A"),
    ("C", "E"),
    ("C", "F"),
    ("G", "A"),
    ("G", "C"),
    ("G", "E"),
    ("D", "C"),
    ("D", "E"),
    ("D", "G"),
    ("A", "G"),
    ("E", "G"),
    ("E", "D"),
    ("F", "G"),
    ("Bb", "C"),
]


@pytest.mark.parametrize("from_key,to_key", RETIRED_KEY_PAIRS)
def test_retired_key_pair_redirects_permanently(
    client: FlaskClient, from_key: str, to_key: str
) -> None:
    slug = f"{seo.key_slug(from_key)}{seo.KEY_PAIR_SLUG_SEPARATOR}{seo.key_slug(to_key)}"
    response = client.get(f"/transpose/{slug}")
    assert response.status_code == 301
    assert response.headers["Location"] == f"/?from={quote_plus(from_key)}&to={quote_plus(to_key)}"


def test_retired_key_pair_redirect_handles_sharp_slugs(client: FlaskClient) -> None:
    response = client.get("/transpose/f-sharp-to-c-sharpm")
    assert response.status_code == 301
    assert response.headers["Location"] == "/?from=F%23&to=C%23m"


def test_retired_key_pair_redirect_target_preselects_both_keys(client: FlaskClient) -> None:
    response = client.get("/transpose/c-to-g", follow_redirects=True)
    assert response.status_code == 200
    body = response.get_data(as_text=True)
    assert 'data-preselect-from="C"' in body
    assert 'data-preselect-to="G"' in body


@pytest.mark.parametrize("slug", ["c-to-h", "zz-to-c", "c", "a-to-b-to-c"])
def test_unknown_key_pair_slug_is_not_found(client: FlaskClient, slug: str) -> None:
    assert client.get(f"/transpose/{slug}").status_code == 404


def test_not_found_page_is_branded(client: FlaskClient) -> None:
    response = client.get("/no-such-page")
    assert response.status_code == 404
    body = response.get_data(as_text=True)
    assert "That page is out of key" in body
    assert "<title>Page not found — Chord Transposer</title>" in body
    assert "tailwind.css" in body


def test_not_found_page_is_noindex_without_canonical(client: FlaskClient) -> None:
    body = client.get("/no-such-page").get_data(as_text=True)
    assert '<meta name="robots" content="noindex, follow" />' in body
    assert 'rel="canonical"' not in body


def test_not_found_page_links_home_and_landing_pages(client: FlaskClient) -> None:
    body = client.get("/no-such-page").get_data(as_text=True)
    assert 'href="/"' in body
    for link in seo.not_found_links():
        assert f'href="{link["path"]}"' in body
    for path in seo.sitemap_paths():
        assert f'href="{path}"' in body


def test_not_found_links_match_landing_pages() -> None:
    paths = {page["path"] for page in seo.landing_pages()}
    assert {link["path"] for link in seo.not_found_links()} <= paths


def test_indexable_pages_keep_canonical_and_index_robots(client: FlaskClient) -> None:
    for path in seo.sitemap_paths():
        body = client.get(path).get_data(as_text=True)
        assert '<meta name="robots" content="index, follow" />' in body
        assert 'rel="canonical"' in body


def test_unknown_key_pair_slug_renders_branded_not_found(client: FlaskClient) -> None:
    response = client.get("/transpose/c-to-h")
    assert response.status_code == 404
    assert "That page is out of key" in response.get_data(as_text=True)


def test_home_page_ignores_unknown_preselect_keys(client: FlaskClient) -> None:
    body = client.get("/?from=H&to=zzz").get_data(as_text=True)
    assert "data-preselect-from" not in body
    assert "data-preselect-to" not in body


def _text_upload(data_bytes: bytes, filename: str = "song.txt") -> dict[str, Any]:
    return {"file": (BytesIO(data_bytes), filename)}


def test_transpose_text_missing_file(client: FlaskClient) -> None:
    response = client.post(
        "/transpose-text",
        data={"current_key": "C", "target_key": "D"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "choose a .txt" in response.get_json()["error"]


def test_transpose_text_wrong_extension(client: FlaskClient) -> None:
    data = _text_upload(b"C G", filename="song.docx")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Only .txt" in response.get_json()["error"]


def test_transpose_text_missing_keys(client: FlaskClient) -> None:
    data = _text_upload(b"C G")
    data.update({"current_key": "", "target_key": ""})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "keys are required" in response.get_json()["error"]


def test_transpose_text_unsupported_format(client: FlaskClient) -> None:
    data = _text_upload(b"C G")
    data.update({"current_key": "C", "target_key": "D", "format": "docx"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Unsupported output format" in response.get_json()["error"]


def test_transpose_text_empty_file(client: FlaskClient) -> None:
    data = _text_upload(b"")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "empty" in response.get_json()["error"]


def test_transpose_text_too_large(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "MAX_UPLOAD_BYTES", 4)
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "too large" in response.get_json()["error"]


def test_transpose_text_invalid_encoding(client: FlaskClient) -> None:
    data = _text_upload(b"\xff\xfe\x00chord")
    data.update({"current_key": "C", "target_key": "D"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "UTF-8" in response.get_json()["error"]


def test_transpose_text_invalid_key(client: FlaskClient) -> None:
    data = _text_upload(b"C G")
    data.update({"current_key": "X", "target_key": "D"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "not a valid key" in response.get_json()["error"]


def test_transpose_text_txt_success(client: FlaskClient) -> None:
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "D", "format": "txt"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.mimetype == "text/plain"
    assert "song_D.txt" in response.headers["Content-Disposition"]
    assert "C->D" in unquote(response.headers["X-Transpose-Changes"])


def test_transpose_text_same_key_reports_no_changes(client: FlaskClient) -> None:
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "C", "format": "txt"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert unquote(response.headers["X-Transpose-Changes"]) == "none"


def test_transpose_text_chordpro_success(client: FlaskClient) -> None:
    data = _text_upload(b"[C]Amazing [G]grace", filename="grace.cho")
    data.update({"current_key": "C", "target_key": "D", "format": "chordpro"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert "grace_D.pro" in response.headers["Content-Disposition"]
    assert response.get_data(as_text=True).startswith("[D]Amazing")


def test_transpose_text_pdf_success(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "convert_docx_to_pdf", lambda docx_bytes: _one_page_pdf())
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "D", "format": "pdf"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert "song_D.pdf" in response.headers["Content-Disposition"]


def test_transpose_text_pdf_conversion_failure(
    client: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _raise(_docx_bytes: bytes) -> NoReturn:
        raise app_module.PdfConversionError("PDF conversion timed out.")

    monkeypatch.setattr(app_module, "convert_docx_to_pdf", _raise)
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "D", "format": "pdf"})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 503
    assert "timed out" in response.get_json()["error"]


def _make_chord_docx(text: str = "C G Am F") -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_binder_no_files(client: FlaskClient) -> None:
    response = client.post(
        "/binder",
        data={"current_key": "C", "target_key": "D"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "at least one" in response.get_json()["error"]


def test_binder_too_many_files(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "MAX_BINDER_FILES", 1)
    data = {
        "files": [(BytesIO(_make_chord_docx()), "a.docx"), (BytesIO(_make_chord_docx()), "b.docx")],
        "current_key": "C",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Too many files" in response.get_json()["error"]


def test_binder_missing_keys(client: FlaskClient) -> None:
    data = {"files": [(BytesIO(_make_chord_docx()), "a.docx")], "current_key": "", "target_key": ""}
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "keys are required" in response.get_json()["error"]


def test_binder_empty_member(client: FlaskClient) -> None:
    data = {
        "files": [(BytesIO(b""), "empty.docx")],
        "current_key": "C",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "empty" in response.get_json()["error"]


def test_binder_total_too_large(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "MAX_BINDER_TOTAL_BYTES", 5)
    data = {
        "files": [(BytesIO(_make_chord_docx()), "a.docx")],
        "current_key": "C",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "combined maximum" in response.get_json()["error"]


def test_binder_invalid_key(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "convert_docx_to_pdf", lambda docx_bytes: _one_page_pdf())
    data = {
        "files": [(BytesIO(_make_chord_docx()), "a.docx")],
        "current_key": "X",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "not a valid key" in response.get_json()["error"]


def test_binder_pdf_conversion_failure(
    client: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _raise(_docx_bytes: bytes) -> NoReturn:
        raise app_module.PdfConversionError("Could not reach the PDF conversion service.")

    monkeypatch.setattr(app_module, "convert_docx_to_pdf", _raise)
    data = {
        "files": [(BytesIO(_make_chord_docx()), "a.docx")],
        "current_key": "C",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 503
    assert "conversion service" in response.get_json()["error"]


def test_binder_success_merges_pages(client: FlaskClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "convert_docx_to_pdf", lambda docx_bytes: _one_page_pdf())
    data = {
        "files": [
            (BytesIO(_make_chord_docx("C G")), "a.docx"),
            (BytesIO(_make_chord_docx("F C")), "b.docx"),
        ],
        "current_key": "C",
        "target_key": "D",
    }
    response = client.post("/binder", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert len(PdfReader(BytesIO(response.data)).pages) == 2
    assert "chord_binder_D.pdf" in response.headers["Content-Disposition"]


def test_detect_key_missing_file(client: FlaskClient) -> None:
    response = client.post("/detect-key", data={}, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "choose a .docx" in response.get_json()["error"]


def test_detect_key_wrong_extension(client: FlaskClient) -> None:
    data = _upload(b"C G Am F", filename="song.txt")
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "Only .docx" in response.get_json()["error"]


def test_detect_key_empty_file(client: FlaskClient) -> None:
    data = _upload(b"", filename="empty.docx")
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "empty" in response.get_json()["error"]


def test_detect_key_too_large(client: FlaskClient) -> None:
    data = _upload(b"x" * (10 * 1024 * 1024 + 1))
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "too large" in response.get_json()["error"]


def test_detect_key_invalid_docx(client: FlaskClient) -> None:
    data = _upload(b"this is not a real docx", filename="bad.docx")
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "valid .docx" in response.get_json()["error"]


def test_detect_key_success(client: FlaskClient) -> None:
    document = Document()
    document.add_paragraph("C F G7 C")
    document.add_paragraph("These are the lyrics")
    buffer = BytesIO()
    document.save(buffer)
    data = _upload(buffer.getvalue())
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    result = response.get_json()
    assert result["key"] == "C"
    assert result["ambiguous"] is False
    assert result["chords"] == 4


def test_detect_key_too_few_chords(client: FlaskClient) -> None:
    document = Document()
    document.add_paragraph("Only lyrics in this one")
    buffer = BytesIO()
    document.save(buffer)
    data = _upload(buffer.getvalue())
    response = client.post("/detect-key", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.get_json()["key"] is None


def _edits(*pairs: tuple[str, str]) -> str:
    return json.dumps([{"source": source, "text": text} for source, text in pairs])


def test_transpose_applies_preview_edits(client: FlaskClient) -> None:
    data = _upload(_chord_docx_bytes())
    data.update(
        {
            "current_key": "C",
            "target_key": "D",
            "edits": _edits(("C G Am F", "D A Bm7 G")),
        }
    )
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    assert paragraphs == ["D A Bm7 G", "These are the lyrics"]


@pytest.mark.parametrize(
    "edits",
    ["not json", '{"source": "a"}', "[1]", '[{"source": 1, "text": "x"}]', '[{"source": "a"}]'],
)
def test_transpose_rejects_malformed_edits(client: FlaskClient, edits: str) -> None:
    data = _upload(_chord_docx_bytes())
    data.update({"current_key": "C", "target_key": "D", "edits": edits})
    response = client.post("/transpose", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert response.get_json()["error"] == app_module.EDITS_ERROR


def test_transpose_text_applies_preview_edits(client: FlaskClient) -> None:
    data = _text_upload(b"C G Am F\nlyrics")
    data.update(
        {
            "current_key": "C",
            "target_key": "D",
            "format": "txt",
            "edits": _edits(("C G Am F", "D A Bm7 G")),
        }
    )
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.get_data(as_text=True) == "D A Bm7 G\nlyrics"


def test_transpose_text_rejects_malformed_edits(client: FlaskClient) -> None:
    data = _text_upload(b"C G Am F")
    data.update({"current_key": "C", "target_key": "D", "format": "txt", "edits": "[["})
    response = client.post("/transpose-text", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    assert response.get_json()["error"] == app_module.EDITS_ERROR
